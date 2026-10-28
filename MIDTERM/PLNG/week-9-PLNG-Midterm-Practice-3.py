# PLNG - Midterm Practice #3 (Programs 31 - 50)


from random import randint  # randint gives a random whole number


# ---------- Helper functions ----------

def ask_int_list(message):
    # ask for numbers separated by commas, e.g. 3,1,2
    text = input(message)                        # read the text
    return [int(x) for x in text.split(",")]     # split by comma and make each one an integer


def ask_text_list(message):
    # ask for words separated by commas, e.g. Monday,Tuesday
    text = input(message)                        # read the text
    return [x.strip() for x in text.split(",")]  # split by comma and remove extra spaces


# ---------- Program 31: read records from a database ----------

def program31():
    import pymysql.cursors                       # module to talk to MySQL
    db_name = input("Database name (example: students): ")
    user = input("Database user (example: root): ")
    password = input("Database password (can be empty): ")
    table = input("Table name (example: users): ")

    # connect to the database
    connection = pymysql.connect(host='localhost', user=user, password=password,
                                 db=db_name, charset='utf8mb4',
                                 cursorclass=pymysql.cursors.DictCursor)
    try:
        with connection.cursor() as cursor:      # cursor runs the SQL commands
            sql = "SELECT `id`, `firstname`, `lastname` FROM `" + table + "`"
            cursor.execute(sql)                  # run the SQL command
            for row in cursor.fetchall():        # go through every record
                print("firstname : " + str(row["firstname"]))
                print("lastname : " + str(row["lastname"]))
    finally:
        connection.close()                       # always close the connection


# ---------- Program 32: simple Tkinter form ----------

def program32():
    from tkinter import Tk
    title = input("Window title: ")
    window = Tk()                                # create the window
    window.title(title)                          # set the window title
    window.geometry("300x200+10+20")             # width x height + x position + y position
    window.mainloop()                            # keep the window open


# ---------- Program 33: form with a label and an entry ----------

def program33():
    from tkinter import Tk, Frame, Label, Entry
    label_text = input("Label text (example: Enter your name): ")
    window = Tk()
    window.title("Form Entry")
    window.geometry("400x300")
    application = Frame(window)                  # frame holds the widgets
    application.grid()                           # use grid layout
    L1 = Label(application, text=label_text)     # text label
    L1.grid(padx=110, pady=10)                   # place the label
    E1 = Entry(application, bd=2)                # text box
    E1.grid(padx=110, pady=3)                    # place the text box
    window.mainloop()


# ---------- Program 34: Tkinter ListBox ----------

def program34():
    from tkinter import Tk, Frame, Listbox
    items = ask_text_list("Enter list items separated by commas: ")
    window = Tk()
    window.title("ListBox")
    window.geometry("400x300")
    application = Frame(window)
    application.grid()
    Lb1 = Listbox(application)                   # create the list box
    for number, item in enumerate(items, start=1):
        Lb1.insert(number, item)                 # add each item to the list box
    Lb1.grid(padx=110, pady=10)
    window.mainloop()


# ---------- Program 35: registration form saved to Excel ----------

def program35():
    from openpyxl import Workbook, load_workbook  # modules for Excel files
    from tkinter import Tk, Label, Entry, Button, END
    import os

    path = input("Excel file path (example: excel.xlsx): ")
    if os.path.exists(path):                     # file already exists
        wb = load_workbook(path)                 # open it
    else:
        wb = Workbook()                          # make a new Excel file
    sheet = wb.active                            # use the active sheet

    names = ["Name", "Course", "Semester", "Form Number",
             "Contact Number", "Email id", "Address"]
    widths = [30, 10, 10, 20, 20, 40, 50]
    for i in range(7):
        sheet.column_dimensions[chr(65 + i)].width = widths[i]  # set column width (A, B, C...)
        sheet.cell(row=1, column=i + 1).value = names[i]        # write the header

    root = Tk()
    root.configure(background='light green')     # background color
    root.title("registration form")
    root.geometry("500x300")
    Label(root, text="Form", bg="light green").grid(row=0, column=1)

    fields = []                                  # list of text boxes
    for i in range(7):
        Label(root, text=names[i], bg="light green").grid(row=i + 1, column=0)
        entry = Entry(root)
        entry.grid(row=i + 1, column=1, ipadx="100")
        fields.append(entry)

    # when Enter is pressed, move the cursor to the next box
    for i in range(6):
        fields[i].bind("<Return>", lambda event, n=i + 1: fields[n].focus_set())

    def clear():
        for f in fields:
            f.delete(0, END)                     # empty each text box

    def insert():
        values = [f.get() for f in fields]       # get the text from each box
        if all(v == "" for v in values):         # nothing was typed
            print("empty input")
        else:
            row = sheet.max_row + 1              # next empty row
            for i in range(7):
                sheet.cell(row=row, column=i + 1).value = values[i]
            wb.save(path)                        # save the Excel file
            fields[0].focus_set()                # cursor back to the first box
            clear()

    Button(root, text="Submit", fg="Black", bg="Red", command=insert).grid(row=8, column=1)
    root.mainloop()


# ---------- Program 36: number guessing game ----------

def program36():
    rand = randint(1, 100)                       # secret number
    counter = 0                                  # number of guesses
    while True:
        counter += 1
        num = int(input("Enter values between 1 and 100 (0 exit): "))
        if num == 0:
            print("You canceled the game.")
            break
        elif num < rand:
            print("Enter a Higher Number.")
        elif num > rand:
            print("Enter a Lower Number.")
        else:
            print("Randomly selected number {0}!".format(rand))
            print("Your guess number {0}".format(counter))
            break                                # correct guess, stop the game


# ---------- Program 37: favorite fruits ----------

def program37():
    favorite_fruits = ask_text_list("Enter your favorite fruits separated by commas: ")
    print("My favourite fruits {}".format(favorite_fruits))


# ---------- Program 38: mixed list ----------

def program38():
    pi = float(input("Number of pi (example: 3.14): "))
    inch = float(input("1 inch in cm (example: 2.54): "))
    micro = input("Abbreviation of microprocessor (example: CPU): ")
    os_name = input("Your operating system: ")
    my_list = [pi, inch, micro, os_name]         # list with different types
    print(my_list)


# ---------- Program 39: weekdays and an index ----------

def program39():
    weekdays = ask_text_list("Enter the weekdays separated by commas: ")
    index = int(input("Which index do you want to print? (starts at 0): "))
    if index < len(weekdays):                    # check the index exists
        print(weekdays[index])                   # print the element at that index
    else:
        print("This index is not in the list.")


# ---------- Program 40: change list values ----------

def program40():
    my_list = ask_int_list("Enter 3 numbers separated by commas (example 3,1,2): ")
    print("Old list:", my_list)
    new_values = ask_int_list("Enter 3 NEW numbers separated by commas (example 1,1,2): ")
    for i in range(len(my_list)):
        my_list[i] = new_values[i]               # replace each value
    print("New list:", my_list)


# ---------- Program 41: check if a day is in the list ----------

def program41():
    weekdays = ask_text_list("Enter the weekdays separated by commas: ")
    while True:
        day = input("Enter a day to check (0 to stop): ")
        if day == "0":
            break
        print(day in weekdays)                   # True if it is in the list, else False


# ---------- Program 42: even numbers ----------

def program42():
    list1 = ask_int_list("Enter numbers separated by commas: ")
    for num in list1:
        if num % 2 == 0:                         # remainder 0 means even
            print(num, end=" ")                  # print on the same line
    print()                                      # new line at the end


# ---------- Program 43: combine two lists, print numbers divisible by 4 ----------

def program43():
    start = int(input("Start number (example: 3): "))
    end = int(input("End number (example: 18): "))
    numbers = list(range(start, end + 1))        # numbers from start to end
    numbers2 = ask_int_list("Enter the second list separated by commas (example 21,22,23): ")
    combine = numbers + numbers2                 # join the two lists
    for number in combine:
        if number % 4 == 0:                      # divisible by 4
            print(number)


# ---------- Program 44: while loop with classes ----------

def program44():
    total = int(input("How many classes? (example: 12): "))
    i = 1                                        # counter starts at 1
    while i <= total:
        print(" {}. class".format(i))
        i += 1                                   # add 1 each time


# ---------- Program 45: sum of digits ----------

def program45():
    number = input("Enter a number: ")           # keep it as text so we can loop each digit
    total = 0
    for digit in number:
        total += int(digit)                      # add each digit
    print("The sum of the digits of the number:", total)


# ---------- Program 46: multiply digits (except zero) ----------

def program46():
    number = input("Enter a number: ")
    multiplication = 1                           # start at 1 (not 0)
    for digit in number:
        if int(digit) != 0:                      # skip the zeros
            multiplication *= int(digit)
    print("Multiplication of the digits:", multiplication)


# ---------- Program 47: print name with a for loop ----------

def program47():
    text = input("Enter a name: ")
    times = int(input("How many times? (example: 10): "))
    for i in range(times):                       # repeat 'times' times
        print(text)


# ---------- Program 48: print name with a while loop ----------

def program48():
    text = input("Enter a name: ")
    times = int(input("How many times? (example: 10): "))
    i = 0
    while i < times:
        i += 1
        print(text)


# ---------- Program 49: mean with sum() and len() ----------

def program49():
    numbers = ask_int_list("Enter numbers separated by commas: ")
    total = sum(numbers)                         # add all numbers
    piece = len(numbers)                         # count the numbers
    print("Mean:", total / piece)


# ---------- Program 50: sum until 0 is entered ----------

def program50():
    total = 0
    while True:
        num = float(input("Enter a number (0 to stop): "))
        if num == 0:                             # stop when the user types 0
            break
        total += num
    print("The sum of the numbers:", total)


# ---------- Menu (dictionary: number -> [title, function]) ----------

programs = {
    31: ["Read records from a database", program31],
    32: ["Tkinter form", program32],
    33: ["Form with label and entry", program33],
    34: ["Tkinter ListBox", program34],
    35: ["Registration form (Excel)", program35],
    36: ["Number guessing game", program36],
    37: ["Favorite fruits list", program37],
    38: ["Pi, inch, CPU, OS list", program38],
    39: ["Weekdays and index", program39],
    40: ["Change list values", program40],
    41: ["Check days in a list", program41],
    42: ["Even numbers in a list", program42],
    43: ["Combine lists, divisible by 4", program43],
    44: ["While loop: classes", program44],
    45: ["Sum of digits", program45],
    46: ["Multiply digits (no zero)", program46],
    47: ["Print name (for loop)", program47],
    48: ["Print name (while loop)", program48],
    49: ["Mean with sum() and len()", program49],
    50: ["Sum until 0", program50],
}


def main():
    while True:
        print("\n----------     PLNG Midterm Practice 3     ----------")
        for number, info in programs.items():
            print(number, "-", info[0])          # show the menu
        try:
            choice = int(input("Choose a program (31-50): "))
            if choice in programs:
                programs[choice][1]()            # run the chosen program
            else:
                print("Please choose a number between 31 and 50.")
        except Exception as error:               # show the error but do not close the menu
            print("Error:", error)

        again = input("\nDo you want to continue? Y/N: ").strip().upper()
        if again != "Y":                         # anything except Y ends the program
            print("Goodbye!")
            break


# run main() only when this file is started directly
if __name__ == "__main__":
    main()